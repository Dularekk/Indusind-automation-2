import os
import re
import threading
from pathlib import Path
from datetime import datetime
from flask import Flask, jsonify, request, send_file, render_template_string
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright

app = Flask(__name__)
URL = "https://transact.indusindinsurance.com/insurance/self-help/online-health-card.aspx"
BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = BASE_DIR / "output.xlsx"

lock = threading.RLock()
resume_event = threading.Event()
resume_event.set()
state = {
    "running": False, "paused": False, "stop_requested": False,
    "start": "", "end": "", "current": "", "total": 0,
    "processed": 0, "success": 0, "failed": 0,
    "status": "Ready", "message": "Ready", "logs": []
}
worker_thread = None


def log(message):
    with lock:
        state["message"] = message
        state["logs"].append(datetime.now().strftime("%H:%M:%S") + " - " + message)
        state["logs"] = state["logs"][-80:]


def initialize_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Policy No.", "Customer Name", "Expiry Date", "Status"])
    wb.save(OUTPUT_FILE)
    wb.close()


def save_result(policy, name="", expiry="", status="FAILED"):
    with lock:
        if OUTPUT_FILE.exists():
            wb = load_workbook(OUTPUT_FILE)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.append(["Policy No.", "Customer Name", "Expiry Date", "Status"])
        ws.append([str(policy), name, expiry, status])
        wb.save(OUTPUT_FILE)
        wb.close()


def visible_policy_field(page):
    selectors = [
        'input[placeholder*="policy" i]', 'input[name*="policy" i]',
        'input[id*="policy" i]', 'input[aria-label*="policy" i]'
    ]
    for selector in selectors:
        try:
            loc = page.locator(selector).first
            if loc.is_visible(timeout=1000) and loc.is_enabled():
                return loc
        except Exception:
            pass
    try:
        fields = page.locator("input:visible")
        for i in range(fields.count()):
            loc = fields.nth(i)
            typ = (loc.get_attribute("type") or "text").lower()
            if typ in ("text", "number", "search", "tel") and loc.is_enabled():
                return loc
    except Exception:
        pass
    raise RuntimeError("Visible policy number field nahi mila; website layout verify karna hoga.")


def extract_details(page):
    text = page.locator("body").inner_text(timeout=10000)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    name = ""
    expiry = ""
    name_labels = r"customer\s*name|member\s*name|insured\s*name"
    expiry_labels = r"expiry\s*date|policy\s*expiry\s*date|valid\s*(?:till|upto|up\s*to)"
    for i, line in enumerate(lines):
        if not name and re.search(name_labels, line, re.I):
            m = re.search(r"(?:" + name_labels + r")\s*[:\-]?\s*(.+)$", line, re.I)
            if m and m.group(1).strip().lower() not in ("name", ""):
                name = m.group(1).strip()
            elif i + 1 < len(lines):
                name = lines[i + 1].strip()
        if not expiry and re.search(expiry_labels, line, re.I):
            m = re.search(r"(?:" + expiry_labels + r")\s*[:\-]?\s*(.+)$", line, re.I)
            if m and m.group(1).strip().lower() not in ("date", ""):
                expiry = m.group(1).strip()
            elif i + 1 < len(lines):
                expiry = lines[i + 1].strip()
    return name, expiry


def process_policy(page, policy):
    page.goto(URL, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_timeout(1200)
    field = visible_policy_field(page)
    field.fill(str(policy))
    buttons = page.get_by_role("button", name=re.compile(r"confirm|search|submit|go", re.I))
    if buttons.count():
        buttons.first.click(timeout=15000)
    else:
        confirm = page.get_by_text(re.compile(r"^\s*(confirm|search|go)\s*$", re.I)).first
        if confirm.count() == 0:
            raise RuntimeError("Confirm/Search button nahi mila.")
        confirm.click(timeout=15000)
    page.wait_for_timeout(2200)
    name, expiry = extract_details(page)
    if not name and not expiry:
        raise RuntimeError("Result details nahi mile; policy not found ya page layout badla hai.")
    return name, expiry


def worker(start_at):
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
            page = browser.new_page()
            try:
                policy_num = start_at
                while policy_num <= int(state["end"]):
                    resume_event.wait()
                    with lock:
                        if state["stop_requested"]:
                            state["stop_requested"] = False
                            state["running"] = False
                            state["paused"] = True
                            state["status"] = "Paused"
                            log("Paused at policy " + str(policy_num))
                            return
                        state["current"] = str(policy_num)
                    policy = str(policy_num)
                    log("Searching policy " + policy)
                    try:
                        name, expiry = process_policy(page, policy)
                        save_result(policy, name, expiry, "SUCCESS")
                        with lock:
                            state["success"] += 1
                    except Exception as exc:
                        save_result(policy, "", "", "FAILED")
                        with lock:
                            state["failed"] += 1
                        log(policy + " FAILED: " + str(exc)[:180])
                    with lock:
                        state["processed"] += 1
                    policy_num += 1
                with lock:
                    state["running"] = False
                    state["paused"] = False
                    state["status"] = "Completed"
                log("Batch completed")
            finally:
                browser.close()
    except Exception as exc:
        with lock:
            state["running"] = False
            state["status"] = "Error"
        log("Worker error: " + str(exc)[:250])


HTML = r'''<!doctype html>
<html><head><meta name="viewport" content="width=device-width, initial-scale=1">
<title>IndusInd Automation 2</title>
<style>
body{font-family:Arial,sans-serif;background:#f2f4f8;color:#182230;margin:0}main{max-width:620px;margin:auto;padding:16px}
.card{background:white;border-radius:12px;padding:18px;margin:12px 0;box-shadow:0 2px 8px #0001}
input,button{width:100%;box-sizing:border-box;padding:13px;margin:6px 0;border-radius:8px;font-size:16px}
input{border:1px solid #c6ccd5}button{border:0;background:#1756a9;color:white;font-weight:bold}.stop{background:#b42318}.resume{background:#137547}
.stats{display:grid;grid-template-columns:1fr 1fr;gap:10px}.stat{background:#eef3ff;padding:12px;border-radius:8px}
small{color:#555}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:12px;max-height:220px;overflow:auto}
a{display:block;text-align:center;padding:12px}
</style></head><body><main><h2>IndusInd Automation 2</h2>
<div class="card"><form id="form"><label>Start policy number</label><input id="start" inputmode="numeric" pattern="[0-9]+" required>
<label>End policy number</label><input id="end" inputmode="numeric" pattern="[0-9]+" required>
<button type="submit">START NEW RANGE</button></form>
<button class="stop" onclick="action('stop')">STOP / PAUSE</button><button class="resume" onclick="action('resume')">RESUME</button>
<button onclick="action('reset')">RESET FOR NEW RANGE</button></div>
<div class="card"><div><b>Status:</b> <span id="status">Ready</span></div><div><b>Current policy:</b> <span id="current">-</span></div>
<div class="stats"><div class="stat">Progress<br><b id="processed">0</b> / <b id="total">0</b></div><div class="stat">Success<br><b id="success">0</b></div><div class="stat">Failed<br><b id="failed">0</b></div></div>
<p id="message"></p><a href="/download">Download output.xlsx</a><b>Recent activity</b><pre id="logs"></pre></div>
<small>Run one policy first to confirm the website selectors and extracted fields.</small></main>
<script>
const el=id=>document.getElementById(id);
el('form').addEventListener('submit',async e=>{e.preventDefault();let start=el('start').value.trim(),end=el('end').value.trim();if(!/^\d+$/.test(start)||!/^\d+$/.test(end)){alert('Only digits allowed');return}let r=await fetch('/start',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({start,end})});let d=await r.json();alert(d.message);refresh()});
async function action(a){let r=await fetch('/'+a,{method:'POST'});let d=await r.json();alert(d.message);refresh()}
async function refresh(){try{let r=await fetch('/status'),s=await r.json();['status','current','processed','total','success','failed','message'].forEach(k=>el(k).textContent=s[k]??'');el('logs').textContent=(s.logs||[]).slice(-20).join('\n')}catch(e){}}
setInterval(refresh,2000);refresh();
</script></body></html>'''


@app.get("/")
def home():
    return render_template_string(HTML)


@app.get("/status")
def status():
    with lock:
        return jsonify(dict(state))


@app.post("/start")
def start():
    global worker_thread
    data = request.get_json(silent=True) or {}
    start_text = str(data.get("start", "")).strip()
    end_text = str(data.get("end", "")).strip()
    if not re.fullmatch(r"\d+", start_text) or not re.fullmatch(r"\d+", end_text):
        return jsonify(message="Valid policy numbers enter karo."), 400
    start_num, end_num = int(start_text), int(end_text)
    if end_num < start_num:
        return jsonify(message="End number start se chhota nahi ho sakta."), 400
    if end_num - start_num > 100000:
        return jsonify(message="Range bahut bada hai; chhota batch chalao."), 400
    with lock:
        if state["running"]:
            return jsonify(message="Batch already running hai."), 409
        initialize_workbook()
        state.update({"running": True, "paused": False, "stop_requested": False,
            "start": str(start_num), "end": str(end_num), "current": str(start_num),
            "total": end_num-start_num+1, "processed": 0, "success": 0, "failed": 0,
            "status": "Running", "message": "Starting batch", "logs": []})
        resume_event.set()
        worker_thread = threading.Thread(target=worker, args=(start_num,), daemon=True)
        worker_thread.start()
    return jsonify(message="Batch started. Pehle chhota test karo.")


@app.post("/stop")
def stop():
    with lock:
        if not state["running"]:
            return jsonify(message="Koi active batch nahi hai.")
        state["stop_requested"] = True
        state["status"] = "Stopping"
    return jsonify(message="Stop request bhej di; current policy ke baad pause hoga.")


@app.post("/resume")
def resume():
    global worker_thread
    with lock:
        if state["running"]:
            return jsonify(message="Batch already running hai.")
        if not state["paused"]:
            return jsonify(message="Resume karne ke liye paused batch nahi hai."), 400
        next_policy = int(state["current"])
        state["running"] = True
        state["paused"] = False
        state["status"] = "Running"
        state["stop_requested"] = False
        resume_event.set()
        worker_thread = threading.Thread(target=worker, args=(next_policy,), daemon=True)
        worker_thread.start()
    return jsonify(message="Batch resumed.")


@app.post("/reset")
def reset():
    with lock:
        if state["running"]:
            return jsonify(message="Pehle batch STOP karo."), 409
        state.update({"running": False, "paused": False, "stop_requested": False,
            "start": "", "end": "", "current": "", "total": 0,
            "processed": 0, "success": 0, "failed": 0, "status": "Ready",
            "message": "Ready", "logs": []})
    return jsonify(message="Dashboard reset ho gaya. Naya range start kar sakte ho.")


@app.get("/download")
def download():
    if not OUTPUT_FILE.exists():
        initialize_workbook()
    return send_file(OUTPUT_FILE, as_attachment=True, download_name="output.xlsx")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0", port=port)
